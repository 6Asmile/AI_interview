import json
import hashlib
import io
from pathlib import Path
import time
from django.core.management.base import BaseCommand, CommandError


class Command(BaseCommand):
    help = '只核对本轮合成账号的保存、视频完成与会话归属，不调用 AI。'

    def add_arguments(self, parser):
        parser.add_argument('--result', required=True)

    def handle(self, **options):
        from users.models import User
        from resumes.models import ResumeDraft, ResumeArtifact
        from interviews.models import InterviewSession
        from core.models import AsyncOperation, OperationDispatchOutbox
        from video_uploads.models import FileUploadTask, VideoTranscodeTask
        from video_uploads.services import FFmpegService
        path = Path(options['result'])
        result = json.loads(path.read_text(encoding='utf-8'))
        user = User.objects.get(pk=result['user_id'])
        if not user.username.startswith('evaluation-') or not user.email.endswith('@example.invalid'):
            raise CommandError('只允许验证合成验收账号')
        rows = result['rows']
        draft = ResumeDraft.objects.filter(resume_id=result['resume_id'], resume__user=user).first()
        for row in rows:
            if row['name'] in ('Studio 自动保存', 'Studio 连续编辑与自动保存'):
                persisted = bool(draft and draft.resume_json.get('basics', {}).get('summary') == row.pop('expected_summary', None))
                row.update(status='passed' if persisted else 'failed', persisted=persisted)
            if row['name'] == '保存后服务端权威预览可见' and row['status'] == 'passed':
                artifact = ResumeArtifact.objects.get(pk=row['artifact_id'], resume_id=result['resume_id'], resume__user=user)
                # 导出会创建新版本并推进草稿 ETag；核对预览显示当时的
                # ETag 与冻结内容，不要求它等于导出完成后的新 ETag。
                matches = artifact.draft_etag == row['draft_etag'] and artifact.preview_input == draft.resume_json
                row.update(latest_draft_preview=matches, status='passed' if matches else 'failed')
            if row['name'].startswith('浏览器导出并下载 ') and row['status'] == 'passed':
                artifact = ResumeArtifact.objects.select_related('asset', 'content_version').get(
                    pk=row['artifact_id'], resume_id=result['resume_id'], resume__user=user)
                with artifact.asset.file.open('rb') as file:
                    content = file.read()
                expected = draft.resume_json.get('basics', {}).get('summary', '')
                try:
                    if artifact.format == 'json':
                        text = json.loads(content).get('basics', {}).get('summary', '')
                    elif artifact.format == 'docx':
                        from docx import Document
                        document = Document(io.BytesIO(content))
                        text = '\n'.join(p.text for p in document.paragraphs)
                        text += '\n'.join(cell.text for table in document.tables for line in table.rows for cell in line.cells)
                    else:
                        from pypdf import PdfReader
                        text = '\n'.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(content)).pages)
                    preserved = bool(expected and ''.join(expected.split()) in ''.join(text.split()))
                    checksum_valid = hashlib.sha256(content).hexdigest() == artifact.asset.checksum_sha256
                    operation_valid = AsyncOperation.objects.filter(pk=row['operation_id'], user=user, status='succeeded', result_id=str(artifact.pk)).exists()
                    operation = AsyncOperation.objects.get(pk=row['operation_id'], user=user)
                    dispatch = OperationDispatchOutbox.objects.filter(operation=operation).order_by('created_at').first()
                    def milliseconds(start, end):
                        return round((end - start).total_seconds() * 1000, 3) if start and end else None
                    row['durable_timings_ms'] = {
                        'creation_to_claim': milliseconds(operation.created_at, operation.started_at),
                        'claim_to_completion': milliseconds(operation.started_at, operation.completed_at),
                        'outbox_to_confirm': milliseconds(dispatch.created_at, dispatch.published_at) if dispatch else None,
                    }
                    row.update(content_preserved=preserved, checksum_valid=checksum_valid, operation_succeeded=operation_valid,
                        status='passed' if preserved and checksum_valid and operation_valid else 'failed')
                except Exception as error:
                    row.update(status='failed', content_verification_error=type(error).__name__)
        if any(row['name'] == '真实浏览器录像分片上传受理' and row['status'] == 'passed' for row in rows):
            started = time.perf_counter()
            deadline = time.monotonic() + 60
            operation = None
            while time.monotonic() < deadline:
                operation = AsyncOperation.objects.filter(user=user, operation_type='media.video_process').first()
                if operation and operation.status in ('succeeded', 'failed', 'canceled'):
                    break
                time.sleep(.5)
            upload = FileUploadTask.objects.filter(user=user).first()
            transcode = VideoTranscodeTask.objects.filter(user=user).first()
            info = FFmpegService().get_video_info(transcode.transcoded_file) if transcode and transcode.status == 'completed' else None
            processed = bool(operation and operation.status == 'succeeded' and upload and upload.status == 'merged' and upload.uploaded_chunks == upload.total_chunks and info)
            session = InterviewSession.objects.get(pk=result['session_id'], user=user)
            rows.append({'name': '真实视频 Operation 合并转码及可解码结果', 'status': 'passed' if processed else 'failed',
                'operation_id': str(operation.pk) if operation else None, 'upload_id': str(upload.pk) if upload else None,
                'wait_ms': round((time.perf_counter() - started) * 1000, 2),
                'operation_duration_ms': round((operation.completed_at - operation.created_at).total_seconds() * 1000, 2) if operation and operation.completed_at else None,
                'video_seconds': info.get('format', {}).get('duration') if info else None})
            rows.append({'name': '视频与面试会话关联', 'status': 'passed' if upload and session.video_upload_task_id == upload.pk else 'failed'})
        result['status'] = 'failed' if any(row['status'] == 'failed' for row in rows) else 'partial'
        path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        self.stdout.write('合成浏览器数据库与媒体结果已核对')
